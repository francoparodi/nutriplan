"""Resoconto di un intervallo di date (PDF / Excel) per medico e paziente."""
import io
from datetime import date, timedelta
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..deps import today
from ..models import Diet, MealLog, PatientProfile, User
from .day import progress
from .diet import diet_to_dict

T = {
    "it": {"title": "Resoconto alimentare", "patient": "Paziente", "period": "Periodo", "doctor": "Medico",
           "generated": "Generato il", "summary": "Riepilogo", "weights": "Andamento del peso", "meals": "Diario dei pasti",
           "date": "Data", "meal": "Pasto", "status": "Stato", "detail": "Cosa è stato mangiato", "note": "Note",
           "done": "Come da dieta", "off_plan": "Fuori dieta", "skipped": "Saltato", "missing": "Non compilato",
           "kg": "Peso (kg)", "avg_adh": "Aderenza media (%)", "n_done": "Pasti come da dieta",
           "n_off": "Pasti fuori dieta", "n_skip": "Pasti saltati", "n_miss": "Pasti non compilati",
           "n_subs": "Sostituzioni usate", "adjusted": "Quantità modificata", "n_adjusted": "Pasti con quantità modificate", "planned": "previsto", "changed": "Modificata", "w_first": "Peso iniziale", "w_last": "Peso finale", "w_delta": "Variazione",
           "birth": "Data di nascita", "height": "Altezza (cm)", "goal": "Peso obiettivo (kg)", "adh": "Aderenza (%)",
           "required": "Pasti previsti", "subs": "Sostituzioni", "sheet_sum": "Riepilogo", "sheet_meals": "Pasti",
           "sheet_w": "Peso", "sheet_days": "Giorni", "extras": "Extra", "page": "Pagina"},
    "en": {"title": "Food diary report", "patient": "Patient", "period": "Period", "doctor": "Doctor",
           "generated": "Generated on", "summary": "Summary", "weights": "Weight trend", "meals": "Meal diary",
           "date": "Date", "meal": "Meal", "status": "Status", "detail": "What was eaten", "note": "Notes",
           "done": "As planned", "off_plan": "Off plan", "skipped": "Skipped", "missing": "Not filled in",
           "kg": "Weight (kg)", "avg_adh": "Average adherence (%)", "n_done": "Meals as planned",
           "n_off": "Meals off plan", "n_skip": "Meals skipped", "n_miss": "Meals not filled in",
           "n_subs": "Substitutions used", "adjusted": "Quantity changed", "n_adjusted": "Meals with changed quantities", "planned": "planned", "changed": "Changed", "w_first": "Starting weight", "w_last": "Latest weight", "w_delta": "Change",
           "birth": "Date of birth", "height": "Height (cm)", "goal": "Target weight (kg)", "adh": "Adherence (%)",
           "required": "Planned meals", "subs": "Substitutions", "sheet_sum": "Summary", "sheet_meals": "Meals",
           "sheet_w": "Weight", "sheet_days": "Days", "extras": "Extras", "page": "Page"},
}


def collect(db: Session, patient: User, a: date, b: date, lang: str = "it") -> dict:
    t = T.get(lang, T["it"])
    prog = progress(db, patient, a, b)
    diets = list(db.scalars(select(Diet).where(Diet.patient_id == patient.id, Diet.status == "active")))
    docs = {d.id: diet_to_dict(db, d) for d in diets}
    ordered = sorted(diets, key=lambda x: (x.start_date or date.min, x.created_at), reverse=True)
    logs: dict = {}
    for l in db.scalars(select(MealLog).where(MealLog.patient_id == patient.id, MealLog.date >= a,
                                              MealLog.date <= b).order_by(MealLog.date)):
        logs.setdefault(l.date, {})[str(l.meal_id)] = l
    rows = []
    cur, tdy = a, today()
    while cur <= b:
        dt = next((x for x in ordered if (x.start_date or date.min) <= cur and (x.end_date is None or x.end_date >= cur)), None)
        dl = dict(logs.get(cur, {}))
        if dt:
            for m in docs[dt.id]["meals"]:
                l = dl.pop(m["id"], None)
                if l:
                    rows.append(_row(cur, m["name"], l, t))
                elif m["required"] and m["kind"] == "meal" and cur < tdy:
                    rows.append({"date": cur, "meal": m["name"], "status": "missing", "status_label": t["missing"],
                                 "detail": "", "note": "", "adjusted": False, "adjustments": []})
        for l in dl.values():  # registrazioni di pasti non piu' in dieta
            rows.append(_row(cur, l.meal_name or "?", l, t))
        cur += timedelta(days=1)
    pr = db.get(PatientProfile, patient.id)
    return {"t": t, "patient": patient, "profile": pr, "a": a, "b": b, "rows": rows, "progress": prog}


def _row(d, name, l: MealLog, t) -> dict:
    parts = []
    if l.status == "done":
        if l.option_label:
            parts.append(l.option_label)
        for s in l.substitutions or []:
            parts.append(f"{s['item_name']} -> {s['name']} {s.get('qty_text', '')}".strip())
        for e in l.extras or []:
            parts.append(f"{e['name']} x{e.get('count', 1)}")
    adj = [f"{a['item_name']}: {a['planned']:g} {a.get('unit', '')} -> {a['actual']:g} {a.get('unit', '')} ({t['planned']}: {a['planned']:g})"
           for a in (l.adjustments or [])]
    if l.free_text:
        parts.append(l.free_text)
    return {"date": d, "meal": name, "status": l.status, "status_label": t.get(l.status, l.status),
            "detail": "\n".join(parts), "note": l.note, "adjusted": bool(adj), "adjustments": adj}


# ------------------------------------------------------------------ Excel
def report_xlsx(data: dict, doctor_name: str = "") -> bytes:
    t, p, prog = data["t"], data["patient"], data["progress"]
    wb = Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F7D5B")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill = hdr_font, hdr_fill
        ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=1)

    ws = wb.active
    ws.title = t["sheet_sum"]
    s = prog["summary"]
    info = [(t["title"], "Nutriplan"), (t["patient"], p.full_name), (t["doctor"], doctor_name),
            (t["period"], f"{data['a']} → {data['b']}")]
    pr = data["profile"]
    if pr:
        info += [(t["birth"], pr.birth_date.isoformat() if pr.birth_date else ""), (t["height"], pr.height_cm),
                 (t["goal"], pr.goal_weight_kg)]
    info += [(t["required"], s["required"]), (t["n_done"], s["done"]), (t["n_off"], s["off_plan"]),
             (t["n_skip"], s["skipped"]), (t["n_miss"], s["missing"]), (t["n_subs"], s["subs"]),
             (t["n_adjusted"], s["adjusted"]),
             (t["avg_adh"], s["avg_adherence"]), (t["w_first"], s["first_weight"]), (t["w_last"], s["last_weight"]),
             (t["w_delta"], s["delta_weight"])]
    for k, v in info:
        ws.append([k, v])
    for r in ws.iter_rows(min_col=1, max_col=1):
        r[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 36

    w1 = wb.create_sheet(t["sheet_meals"])
    header(w1, [t["date"], t["meal"], t["status"], t["detail"], t["adjusted"], t["note"]])
    amber = PatternFill("solid", fgColor="FFE8B0")
    for r in data["rows"]:
        w1.append([r["date"], r["meal"], r["status_label"], r["detail"], "\n".join(r["adjustments"]), r["note"]])
        if r["adjusted"]:
            for c in w1[w1.max_row]:
                c.fill = amber
    for row in w1.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        row[3].alignment = row[4].alignment = row[5].alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([13, 24, 16, 60, 50, 40], 1):
        w1.column_dimensions[get_column_letter(i)].width = w

    w2 = wb.create_sheet(t["sheet_w"])
    header(w2, [t["date"], t["kg"], t["note"]])
    for w in prog["weights"]:
        w2.append([date.fromisoformat(w["date"]), w["kg"], w["note"]])
    for row in w2.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
    for i, w in enumerate([13, 12, 40], 1):
        w2.column_dimensions[get_column_letter(i)].width = w

    w3 = wb.create_sheet(t["sheet_days"])
    header(w3, [t["date"], t["required"], t["n_done"], t["n_off"], t["n_skip"], t["n_miss"], t["subs"], t["n_adjusted"], t["adh"]])
    for d in prog["days"]:
        w3.append([date.fromisoformat(d["date"]), d["required"], d["done"], d["off_plan"], d["skipped"],
                   d["missing"], d["subs"], d["adjusted"], d["adherence"]])
    for row in w3.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
    for i in range(1, 10):
        w3.column_dimensions[get_column_letter(i)].width = 16
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ------------------------------------------------------------------ PDF
def report_pdf(data: dict, doctor_name: str = "") -> bytes:
    t, p, prog = data["t"], data["patient"], data["progress"]
    buf = io.BytesIO()
    pw, _ = A4
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=18 * mm, title=t["title"], author="Nutriplan")
    W = pw - 30 * mm
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=18, alignment=0)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=12,
                        textColor=colors.HexColor("#2F7D5B"))
    sm = ParagraphStyle("sm", parent=ss["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11)
    green = colors.HexColor("#2F7D5B")
    el = [Paragraph(escape(t["title"]), h1)]
    pr = data["profile"]
    meta = [[t["patient"], p.full_name], [t["period"], f"{data['a']:%d/%m/%Y} – {data['b']:%d/%m/%Y}"]]
    if doctor_name:
        meta.append([t["doctor"], doctor_name])
    if pr and pr.birth_date:
        meta.append([t["birth"], f"{pr.birth_date:%d/%m/%Y}"])
    if pr and pr.height_cm:
        meta.append([t["height"], f"{pr.height_cm:g}"])
    if pr and pr.goal_weight_kg:
        meta.append([t["goal"], f"{pr.goal_weight_kg:g}"])
    meta.append([t["generated"], f"{today():%d/%m/%Y}"])
    mt = Table(meta, colWidths=[40 * mm, W - 40 * mm])
    mt.setStyle(TableStyle([("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9), ("FONT", (1, 0), (1, -1), "Helvetica", 9),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    el += [mt, Paragraph(t["summary"], h2)]
    s = prog["summary"]
    fmt = lambda v: "-" if v is None else (f"{v:g}" if isinstance(v, (int, float)) else str(v))
    rows = [[t["required"], fmt(s["required"]), t["avg_adh"], fmt(s["avg_adherence"])],
            [t["n_done"], fmt(s["done"]), t["w_first"], fmt(s["first_weight"])],
            [t["n_off"], fmt(s["off_plan"]), t["w_last"], fmt(s["last_weight"])],
            [t["n_skip"], fmt(s["skipped"]), t["w_delta"], fmt(s["delta_weight"])],
            [t["n_miss"], fmt(s["missing"]), t["n_subs"], fmt(s["subs"])],
            [t["n_adjusted"], fmt(s["adjusted"]), "", ""]]
    st = Table(rows, colWidths=[W * 0.32, W * 0.18, W * 0.32, W * 0.18])
    st.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("FONT", (1, 0), (1, -1), "Helvetica-Bold", 9),
                            ("FONT", (3, 0), (3, -1), "Helvetica-Bold", 9), ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
                            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F6F3")),
                            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1F6F3"))]))
    el.append(st)
    # grafico peso
    ws_ = prog["weights"]
    if len(ws_) >= 2:
        el.append(Paragraph(t["weights"], h2))
        d0 = date.fromisoformat(ws_[0]["date"])
        pts = [((date.fromisoformat(w["date"]) - d0).days, w["kg"]) for w in ws_]
        dr = Drawing(W, 150)
        lp = LinePlot()
        lp.x, lp.y, lp.width, lp.height = 35, 20, W - 50, 115
        lp.data = [pts]
        lp.lines[0].strokeColor, lp.lines[0].strokeWidth = green, 1.8
        lp.lines[0].symbol = None
        lo, hi = min(k for _, k in pts), max(k for _, k in pts)
        pad = max(0.5, (hi - lo) * 0.15)
        lp.yValueAxis.valueMin, lp.yValueAxis.valueMax = lo - pad, hi + pad
        lp.yValueAxis.labels.fontSize = lp.xValueAxis.labels.fontSize = 7
        lp.xValueAxis.valueMin = 0
        lp.xValueAxis.valueMax = max(1, pts[-1][0])
        lp.xValueAxis.labelTextFormat = lambda v, d0=d0: f"{d0 + timedelta(days=int(v)):%d/%m}"
        dr.add(lp)
        el.append(dr)
        wt = Table([[t["date"], t["kg"], t["note"]]] + [[f"{date.fromisoformat(w['date']):%d/%m/%Y}", f"{w['kg']:g}",
                                                         w["note"]] for w in ws_],
                   colWidths=[30 * mm, 25 * mm, W - 55 * mm], repeatRows=1)
        wt.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5), ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
                                ("BACKGROUND", (0, 0), (-1, 0), green), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey)]))
        el.append(wt)
    elif ws_:
        el += [Paragraph(t["weights"], h2), Paragraph(f"{ws_[0]['date']}: {ws_[0]['kg']:g} kg", sm)]
    # diario
    el.append(Paragraph(t["meals"], h2))
    col = {"done": colors.HexColor("#1B7F4B"), "off_plan": colors.HexColor("#B45309"),
           "skipped": colors.HexColor("#6B7280"), "missing": colors.HexColor("#B91C1C")}
    body = [[t["date"], t["meal"], t["status"], t["detail"]]]
    styles = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5), ("BACKGROUND", (0, 0), (-1, 0), green),
              ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
              ("VALIGN", (0, 0), (-1, -1), "TOP")]
    for i, r in enumerate(data["rows"], 1):
        det = escape(r["detail"]).replace("\n", "<br/>")
        if r["adjusted"]:
            det += '<br/><font color="#B45309"><b>' + escape(t["adjusted"]) + ":</b> " + "<br/>".join(escape(a) for a in r["adjustments"]) + "</font>"
        if r["note"]:
            det += f'<br/><i>{escape(r["note"])}</i>'
        body.append([f"{r['date']:%d/%m/%Y}", Paragraph(escape(r["meal"]), sm), r["status_label"], Paragraph(det, sm)])
        styles.append(("TEXTCOLOR", (2, i), (2, i), col.get(r["status"], colors.black)))
        if r["adjusted"]:
            styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF3DC")))
        styles.append(("FONT", (0, i), (-1, i), "Helvetica", 8.5))
    bt = Table(body, colWidths=[22 * mm, 38 * mm, 28 * mm, W - 88 * mm], repeatRows=1)
    bt.setStyle(TableStyle(styles))
    el.append(bt)

    def footer(c, d):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.grey)
        c.drawString(15 * mm, 10 * mm, f"Nutriplan – {p.full_name}")
        c.drawRightString(pw - 15 * mm, 10 * mm, f"{t['page']} {d.page}")
        c.restoreState()

    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
